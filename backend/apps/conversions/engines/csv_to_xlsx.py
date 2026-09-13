"""
CsvToXlsxEngine — converts CSV files to XLSX workbooks.

Uses Python's native csv module for robust parsing and openpyxl for Excel workbook generation.

Features:
  - Auto-detects UTF-8-SIG, UTF-8, and Latin-1 encodings.
  - Preserves Tamil and all international Unicode characters.
  - Preserves leading-zero string values (e.g. "00123", "00045") as text cells.
  - Converts safe numeric values (integers and floats) to Excel numeric types.
  - Applies header formatting (bold text, freeze top row A2, auto-filter).
  - Dynamically calculates column widths clamped between 12 and 50 characters.
  - Handles empty and header-only CSV files gracefully without crashing.
"""

import csv
import logging
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.conversions.engines.base import BaseConversionEngine, ConversionError
from apps.conversions.engines.validators import validate_csv_signature, validate_output_file

logger = logging.getLogger(__name__)


def _parse_cell_value(val: str) -> str | int | float:
    """
    Safely parse cell string into native types or string.

    Rules:
    1. If val starts with '0', has length > 1, and is digits (e.g. '00123'), preserve as str.
    2. If val is integer (e.g. '123', '-45'), convert to int.
    3. If val is float (e.g. '1500.50', '-12.34'), convert to float.
    4. Otherwise, preserve as string (including Unicode, Tamil, identifiers, text).
    """
    if not val:
        return ""

    # Rule 1: Preserve leading-zero values (e.g., postal codes, IDs)
    if len(val) > 1 and val.startswith("0") and val.isdigit():
        return val

    # Rule 2: Integer check
    if val.lstrip("-").isdigit():
        try:
            return int(val)
        except ValueError:
            pass

    # Rule 3: Float check
    if "." in val:
        clean = val.lstrip("-").replace(".", "", 1)
        if clean.isdigit():
            try:
                return float(val)
            except ValueError:
                pass

    return val


def _detect_and_open_csv(input_path: str):
    """
    Open CSV file attempting encodings in order: utf-8-sig, utf-8, latin-1.

    Returns open file handle and encoding name.
    """
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            f = open(input_path, "r", encoding=encoding, newline="")
            # Read first chunk to test validity
            f.read(1024)
            f.seek(0)
            return f, encoding
        except (UnicodeDecodeError, OSError):
            continue

    raise ConversionError("Unable to decode CSV file with supported encodings (UTF-8, Latin-1).")


class CsvToXlsxEngine(BaseConversionEngine):
    """
    Conversion engine for CSV → XLSX using native python csv and openpyxl.
    """

    source_format = "csv"
    target_format = "xlsx"

    def convert(self, input_path: str, output_path: str) -> str | None:
        """
        Convert CSV file to XLSX workbook.

        Parameters
        ----------
        input_path : str
            Absolute path to input CSV file.
        output_path : str
            Absolute destination path for generated XLSX.

        Returns
        -------
        str | None
            None, indicating output_path was used directly.
        """
        # 1. Validate CSV signature
        validate_csv_signature(input_path)

        # 2. Open CSV with encoding fallback
        f = None
        try:
            f, encoding = _detect_and_open_csv(input_path)
            reader = csv.reader(f)

            # 3. Create openpyxl workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            rows = list(reader)

            # 4. Populate rows into worksheet
            if rows:
                for r_idx, row in enumerate(rows, start=1):
                    parsed_row = [_parse_cell_value(cell) for cell in row]
                    ws.append(parsed_row)

                # Header formatting (if at least 1 row exists)
                header_font = Font(bold=True)
                for cell in ws[1]:
                    cell.font = header_font

                # Freeze pane A2 & Autofilter
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
        except ConversionError:
            raise
        except Exception as exc:
            logger.warning("Error processing CSV file %s: %s", input_path, exc)
            raise ConversionError("Failed to parse CSV file. The file may be corrupt or contain invalid characters.") from exc
        finally:
            if f is not None:
                try:
                    f.close()
                except Exception:  # noqa: BLE001
                    pass

        if rows:
            # 5. Calculate and set column widths (clamped between 12 and 50)
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    val = str(cell.value or "")
                    # If multiline, check max line length
                    lines = val.splitlines()
                    for line in lines:
                        if len(line) > max_len:
                            max_len = len(line)

                col_letter = get_column_letter(col[0].column)
                clamped_width = max(min(max_len + 4, 50), 12)
                ws.column_dimensions[col_letter].width = clamped_width

        # 6. Save workbook
        try:
            out_p = Path(output_path)
            out_p.parent.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            wb.close()
        except Exception as exc:
            logger.exception("Failed to save XLSX workbook to %s", output_path)
            raise ConversionError("Failed to save converted XLSX file.") from exc

        # 7. Validate output XLSX
        validate_output_file(output_path)
        try:
            test_wb = openpyxl.load_workbook(output_path, read_only=True)
            test_wb.close()
        except Exception as exc:
            logger.error("Generated XLSX failed load validation: %s", exc)
            raise ConversionError("The generated XLSX file is invalid or corrupted.") from exc

        logger.info("CsvToXlsxEngine: successfully converted CSV to XLSX (%d rows): %s", len(rows), output_path)
        return None
