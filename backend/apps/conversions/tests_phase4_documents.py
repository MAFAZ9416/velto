"""
Phase 4 Document Conversion Engine Test Suite for VELTO Conversion.

Covers:
  - TXT → PDF, TXT → DOCX
  - HTML → PDF, HTML → DOCX
  - Markdown → PDF, Markdown → DOCX
  - Encoding (UTF-8, UTF-8 BOM, Tamil/Unicode)
  - Line ending normalization (CRLF / LF)
  - Security (script stripping, event handlers, unsafe links, local file blocking)
  - Output validation (PDF %PDF, DOCX ZIP word/document.xml)
  - API, session isolation, failure tracking, and temporary file cleanup.
"""

import io
import os
from pathlib import Path
import tempfile
import zipfile

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

import fitz  # PyMuPDF
import docx

from apps.conversions.engines.base import ConversionError
from apps.conversions.engines.document_engine import (
    DocumentDocxRenderer,
    DocumentPdfRenderer,
    DocumentTextService,
    HtmlSanitizationService,
    HtmlToDocxEngine,
    HtmlToPdfEngine,
    MarkdownRenderingService,
    MarkdownToDocxEngine,
    MarkdownToPdfEngine,
    TxtToDocxEngine,
    TxtToPdfEngine,
)
from apps.conversions.engines.registry import engine_registry
from apps.conversions.engines.validators import (
    validate_docx_output,
    validate_html_signature,
    validate_md_signature,
    validate_pdf_output,
    validate_txt_signature,
)
from apps.conversions.formats import (
    FORMAT_DOCX,
    FORMAT_HTML,
    FORMAT_MD,
    FORMAT_PDF,
    FORMAT_TXT,
)
from apps.conversions.models import ConversionJob, JobStatus


class Phase4DocumentEngineTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.client = APIClient()

    def _create_temp_file(self, content: bytes, suffix: str) -> str:
        fpath = os.path.join(self.temp_dir.name, f"test_input{suffix}")
        with open(fpath, "wb") as f:
            f.write(content)
        return fpath

    def _out_path(self, suffix: str) -> str:
        return os.path.join(self.temp_dir.name, f"test_output{suffix}")

    # ── 1. TXT Engine Tests ───────────────────────────────────────────────────

    def test_txt_to_pdf_utf8_success(self):
        content = "Hello World\nThis is a standard UTF-8 text file.".encode("utf-8")
        inp = self._create_temp_file(content, ".txt")
        out = self._out_path(".pdf")
        engine = TxtToPdfEngine()
        result = engine.convert(inp, out)
        self.assertTrue(os.path.exists(result))
        validate_pdf_output(result)

        doc = fitz.open(result)
        text = doc[0].get_text()
        doc.close()
        self.assertIn("Hello World", text)

    def test_txt_to_docx_utf8_success(self):
        content = "Hello World\nThis is a standard UTF-8 text file.".encode("utf-8")
        inp = self._create_temp_file(content, ".txt")
        out = self._out_path(".docx")
        engine = TxtToDocxEngine()
        result = engine.convert(inp, out)
        self.assertTrue(os.path.exists(result))
        validate_docx_output(result)

        d = docx.Document(result)
        full_text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Hello World", full_text)

    def test_txt_utf8_bom_support(self):
        content = "BOM Text\nWith UTF-8 BOM encoding.".encode("utf-8-sig")
        inp = self._create_temp_file(content, ".txt")
        out_pdf = self._out_path("_bom.pdf")
        engine = TxtToPdfEngine()
        result = engine.convert(inp, out_pdf)
        validate_pdf_output(result)

        doc = fitz.open(result)
        text = doc[0].get_text()
        doc.close()
        self.assertIn("BOM Text", text)

    def test_txt_tamil_unicode_support(self):
        tamil_text = "வணக்கம் உலகம்\nதமிழ் மொழி உரை PDF மற்றும் DOCX ஆக மாறுகிறது."
        content = tamil_text.encode("utf-8")
        inp = self._create_temp_file(content, ".txt")

        # Test PDF
        out_pdf = self._out_path("_tamil.pdf")
        TxtToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)
        doc = fitz.open(out_pdf)
        pdf_text = doc[0].get_text()
        doc.close()
        self.assertIn("வணக்கம்", pdf_text)

        # Test DOCX
        out_docx = self._out_path("_tamil.docx")
        TxtToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)
        d = docx.Document(out_docx)
        docx_text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("வணக்கம்", docx_text)

    def test_txt_crlf_and_lf_normalization(self):
        crlf_text = "Line 1\r\nLine 2\r\nLine 3\r\n".encode("utf-8")
        inp = self._create_temp_file(crlf_text, ".txt")
        normalized = DocumentTextService.decode_and_normalize(inp)
        self.assertNotIn("\r", normalized)
        self.assertEqual(normalized, "Line 1\nLine 2\nLine 3\n")

    def test_txt_empty_file_rejected(self):
        inp = self._create_temp_file(b"", ".txt")
        out = self._out_path(".pdf")
        with self.assertRaises(ConversionError):
            TxtToPdfEngine().convert(inp, out)

    def test_txt_long_lines_wrapping(self):
        long_line = "Word " * 200 + "\nEnd of long line."
        inp = self._create_temp_file(long_line.encode("utf-8"), ".txt")
        out_pdf = self._out_path("_long.pdf")
        out_docx = self._out_path("_long.docx")

        TxtToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        TxtToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

    def test_txt_multipage_pdf_generation(self):
        many_lines = "\n".join(f"Paragraph line number {i}" for i in range(300))
        inp = self._create_temp_file(many_lines.encode("utf-8"), ".txt")
        out_pdf = self._out_path("_multi.pdf")
        TxtToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        doc = fitz.open(out_pdf)
        page_count = len(doc)
        doc.close()
        self.assertGreater(page_count, 1)

    # ── 2. HTML Engine Tests & Security ───────────────────────────────────────

    def test_html_to_pdf_headings_paragraphs_lists(self):
        html_doc = """
        <html>
          <body>
            <h1>Main Title</h1>
            <h2>Subtitle</h2>
            <p>This is a paragraph with <b>bold</b> and <i>italic</i> text.</p>
            <ul>
              <li>Item A</li>
              <li>Item B</li>
            </ul>
          </body>
        </html>
        """
        inp = self._create_temp_file(html_doc.encode("utf-8"), ".html")
        out_pdf = self._out_path("_html.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        doc = fitz.open(out_pdf)
        pdf_text = doc[0].get_text()
        doc.close()
        self.assertIn("Main Title", pdf_text)
        self.assertIn("Item A", pdf_text)

    def test_html_to_docx_headings_paragraphs_lists(self):
        html_doc = """
        <html>
          <body>
            <h1>Main Title</h1>
            <p>Paragraph with <b>bold</b> text.</p>
            <ol>
              <li>First</li>
              <li>Second</li>
            </ol>
          </body>
        </html>
        """
        inp = self._create_temp_file(html_doc.encode("utf-8"), ".html")
        out_docx = self._out_path("_html.docx")
        HtmlToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

        d = docx.Document(out_docx)
        full_text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Main Title", full_text)
        self.assertIn("First", full_text)

    def test_html_table_support(self):
        html_table = """
        <table>
          <tr><th>Name</th><th>Role</th></tr>
          <tr><td>Alice</td><td>Admin</td></tr>
          <tr><td>Bob</td><td>User</td></tr>
        </table>
        """
        inp = self._create_temp_file(html_table.encode("utf-8"), ".html")
        out_pdf = self._out_path("_table.pdf")
        out_docx = self._out_path("_table.docx")

        HtmlToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        HtmlToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

    def test_html_security_script_removal(self):
        malicious_html = """
        <div>
          <script>alert('XSS Attack');</script>
          <p>Safe paragraph text</p>
        </div>
        """
        inp = self._create_temp_file(malicious_html.encode("utf-8"), ".html")
        doc_obj = HtmlSanitizationService.parse_html_to_document(malicious_html)
        out_pdf = self._out_path("_script.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)

        doc = fitz.open(out_pdf)
        text = doc[0].get_text()
        doc.close()
        self.assertNotIn("alert", text)
        self.assertNotIn("XSS Attack", text)
        self.assertIn("Safe paragraph text", text)

    def test_html_security_style_iframe_object_embed_meta_removal(self):
        html_payload = """
        <style>body { background: red; }</style>
        <iframe src="http://evil.com"></iframe>
        <object data="malware.exe"></object>
        <embed src="flash.swf"></embed>
        <meta http-equiv="refresh" content="0;url=http://evil.com">
        <p>Clean text content</p>
        """
        inp = self._create_temp_file(html_payload.encode("utf-8"), ".html")
        out_pdf = self._out_path("_tags.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)

        doc = fitz.open(out_pdf)
        text = doc[0].get_text()
        doc.close()
        self.assertNotIn("evil.com", text)
        self.assertNotIn("malware", text)
        self.assertIn("Clean text content", text)

    def test_html_security_event_handler_removal(self):
        html_event = '<p onclick="evilCode()" onerror="badFunc()">Clickable text</p>'
        inp = self._create_temp_file(html_event.encode("utf-8"), ".html")
        out_pdf = self._out_path("_event.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)

        doc = fitz.open(out_pdf)
        text = doc[0].get_text()
        doc.close()
        self.assertNotIn("evilCode", text)
        self.assertIn("Clickable text", text)

    def test_html_security_dangerous_url_blocking(self):
        html_links = """
        <a href="javascript:alert(1)">JS Link</a>
        <a href="file:///etc/passwd">File Link</a>
        <a href="data:text/html,<script>alert(1)</script>">Data Link</a>
        <a href="https://example.com">Safe Link</a>
        """
        doc_obj = HtmlSanitizationService.parse_html_to_document(html_links)
        runs = []
        for b in doc_obj.blocks:
            if hasattr(b, "runs"):
                runs.extend(b.runs)
        
        links = [r.link for r in runs if r.link]
        self.assertNotIn("javascript:alert(1)", links)
        self.assertNotIn("file:///etc/passwd", links)
        self.assertNotIn("data:text/html,<script>alert(1)</script>", links)

    def test_html_security_local_filesystem_blocking(self):
        html_local = '<a href="C:\\Windows\\System32\\cmd.exe">Cmd</a>'
        doc_obj = HtmlSanitizationService.parse_html_to_document(html_local)
        runs = [r for b in doc_obj.blocks if hasattr(b, "runs") for r in b.runs]
        links = [r.link for r in runs if r.link]
        self.assertEqual(links, [])

    def test_html_security_external_resource_blocking(self):
        html_img = '<img src="http://external-tracker.com/pixel.gif" alt="Image Alt Text">'
        inp = self._create_temp_file(html_img.encode("utf-8"), ".html")
        out_pdf = self._out_path("_no_external.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

    def test_malformed_html_graceful_handling(self):
        malformed = "<p>Unclosed paragraph <b>bold <i>italic</p></div>"
        inp = self._create_temp_file(malformed.encode("utf-8"), ".html")
        out_pdf = self._out_path("_malformed.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

    # ── 3. Markdown Engine Tests ──────────────────────────────────────────────

    def test_markdown_to_pdf_conversion(self):
        md_text = """
# Markdown Title

This is a paragraph with **bold** and *italic* text.

* Item 1
* Item 2

> A blockquote

```python
print("Hello World")
```
"""
        inp = self._create_temp_file(md_text.encode("utf-8"), ".md")
        out_pdf = self._out_path("_md.pdf")
        MarkdownToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        doc = fitz.open(out_pdf)
        text = doc[0].get_text()
        doc.close()
        self.assertIn("Markdown Title", text)
        self.assertIn("A blockquote", text)
        self.assertIn("print(\"Hello World\")", text)

    def test_markdown_to_docx_conversion(self):
        md_text = """
# Markdown Title

Paragraph text.

1. First item
2. Second item

```
Code block text
```
"""
        inp = self._create_temp_file(md_text.encode("utf-8"), ".md")
        out_docx = self._out_path("_md.docx")
        MarkdownToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

        d = docx.Document(out_docx)
        full_text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Markdown Title", full_text)
        self.assertIn("Code block text", full_text)

    def test_markdown_code_blocks_and_blockquotes(self):
        md_code = "> Blockquote quote text\n\n```\ncode line 1\ncode line 2\n```"
        inp = self._create_temp_file(md_code.encode("utf-8"), ".md")
        out_pdf = self._out_path("_code.pdf")
        MarkdownToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

    def test_markdown_safe_links(self):
        md_links = "[Safe Link](https://example.com) and [Unsafe Link](javascript:alert(1))"
        doc_obj = MarkdownRenderingService.markdown_to_document(md_links)
        runs = [r for b in doc_obj.blocks if hasattr(b, "runs") for r in b.runs]
        links = [r.link for r in runs if r.link]
        self.assertIn("https://example.com", links)
        self.assertNotIn("javascript:alert(1)", links)

    def test_markdown_raw_html_safety(self):
        md_raw = "Text before\n<script>alert('injected');</script>\nText after"
        inp = self._create_temp_file(md_raw.encode("utf-8"), ".md")
        out_pdf = self._out_path("_rawhtml.pdf")
        MarkdownToPdfEngine().convert(inp, out_pdf)

        doc = fitz.open(out_pdf)
        text = doc[0].get_text()
        doc.close()
        self.assertNotIn("alert", text)
        self.assertIn("Text before", text)

    # ── 4. Validation & Limits Tests ──────────────────────────────────────────

    def test_pdf_output_validation_signature(self):
        out_pdf = self._out_path("_bad_pdf.pdf")
        with open(out_pdf, "wb") as f:
            f.write(b"NOT A PDF FILE DATA")
        with self.assertRaises(ConversionError):
            validate_pdf_output(out_pdf)

    def test_docx_output_zip_validation(self):
        out_docx = self._out_path("_bad_docx.docx")
        with open(out_docx, "wb") as f:
            f.write(b"NOT A ZIP FILE")
        with self.assertRaises(ConversionError):
            validate_docx_output(out_docx)

    def test_missing_input_file_raises_conversion_error(self):
        inp = os.path.join(self.temp_dir.name, "non_existent.txt")
        out = self._out_path(".pdf")
        with self.assertRaises(ConversionError):
            TxtToPdfEngine().convert(inp, out)

    def test_path_traversal_filename_handling(self):
        file_obj = SimpleUploadedFile(
            name="../../etc/passwd.txt",
            content=b"Sample text file content.",
            content_type="text/plain",
        )
        response = self.client.post(
            "/api/conversions/",
            data={
                "file": file_obj,
                "source_format": "txt",
                "target_format": "pdf",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        job_id = response.data["id"]
        job = ConversionJob.objects.get(id=job_id)
        self.assertNotIn("..", job.original_filename)

    def test_oversized_input_file_rejection(self):
        from unittest.mock import patch
        inp = self._create_temp_file(b"Short text content", ".txt")
        with patch("apps.conversions.engines.document_engine.MAX_TEXT_LENGTH", 10):
            with self.assertRaises(ConversionError):
                DocumentTextService.decode_and_normalize(inp)


    def test_unsupported_conversion_pair_rejected_by_api(self):
        file_obj = SimpleUploadedFile(
            name="test.txt",
            content=b"Sample text content.",
            content_type="text/plain",
        )
        response = self.client.post(
            "/api/conversions/",
            data={
                "file": file_obj,
                "source_format": "txt",
                "target_format": "jpg",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("target_format", response.data)

    # ── 5. API Integration, Isolation, & Cleanup Tests ───────────────────────

    def test_txt_to_pdf_api_job_lifecycle(self):
        file_obj = SimpleUploadedFile(
            name="sample.txt",
            content=b"Hello API document conversion!",
            content_type="text/plain",
        )
        response = self.client.post(
            "/api/conversions/",
            data={
                "file": file_obj,
                "source_format": "txt",
                "target_format": "pdf",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        job_id = response.data["id"]
        job = ConversionJob.objects.get(id=job_id)
        self.assertEqual(job.status, JobStatus.COMPLETED)
        self.assertTrue(os.path.exists(job.output_path))
        validate_pdf_output(job.output_path)

        # Verify Download endpoint
        download_resp = self.client.get(f"/api/conversions/{job_id}/download/")
        self.assertEqual(download_resp.status_code, 200)
        self.assertEqual(download_resp["Content-Type"], "application/pdf")

    def test_html_to_docx_api_job_lifecycle(self):
        file_obj = SimpleUploadedFile(
            name="sample.html",
            content=b"<h1>Title</h1><p>HTML paragraph</p>",
            content_type="text/html",
        )
        response = self.client.post(
            "/api/conversions/",
            data={
                "file": file_obj,
                "source_format": "html",
                "target_format": "docx",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        job_id = response.data["id"]
        job = ConversionJob.objects.get(id=job_id)
        self.assertEqual(job.status, JobStatus.COMPLETED)
        self.assertTrue(os.path.exists(job.output_path))
        validate_docx_output(job.output_path)

    def test_markdown_to_pdf_api_job_lifecycle(self):
        file_obj = SimpleUploadedFile(
            name="sample.md",
            content=b"# Markdown Title\n\nContent paragraph.",
            content_type="text/markdown",
        )
        response = self.client.post(
            "/api/conversions/",
            data={
                "file": file_obj,
                "source_format": "md",
                "target_format": "pdf",
            },
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        job_id = response.data["id"]
        job = ConversionJob.objects.get(id=job_id)
        self.assertEqual(job.status, JobStatus.COMPLETED)
        self.assertTrue(os.path.exists(job.output_path))
        validate_pdf_output(job.output_path)

    def test_session_isolation_preserves_privacy(self):
        client_a = APIClient()
        client_b = APIClient()

        file_obj = SimpleUploadedFile("secret.txt", b"Secret data", "text/plain")
        resp_a = client_a.post(
            "/api/conversions/",
            {"file": file_obj, "source_format": "txt", "target_format": "pdf"},
            format="multipart",
        )
        job_id_a = resp_a.data["id"]

        # Client B cannot view or download Client A's job
        resp_b_detail = client_b.get(f"/api/conversions/{job_id_a}/")
        self.assertEqual(resp_b_detail.status_code, 404)

        resp_b_download = client_b.get(f"/api/conversions/{job_id_a}/download/")
        self.assertEqual(resp_b_download.status_code, 404)

    def test_failed_jobs_appear_in_history(self):
        # Trigger failure during engine processing using non-decodable binary bytes
        file_obj = SimpleUploadedFile("bad_encoding.txt", b"\xff\xfe\xfd\xfa", "text/plain")
        response = self.client.post(
            "/api/conversions/",
            {"file": file_obj, "source_format": "txt", "target_format": "pdf"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.data["status"], "failed")

        # Check job status in DB
        failed_jobs = ConversionJob.objects.filter(status=JobStatus.FAILED)
        self.assertGreaterEqual(failed_jobs.count(), 1)
        history_resp = self.client.get("/api/conversions/")
        self.assertEqual(history_resp.status_code, 200)
        results = history_resp.data.get("jobs", history_resp.data.get("results", history_resp.data))
        history_ids = [j["id"] for j in results]
        self.assertIn(str(failed_jobs.first().id), history_ids)


    def test_forced_failure_cleans_up_temporary_files(self):
        """Verify that on engine failure, temporary input & output files are cleaned up."""
        file_obj = SimpleUploadedFile("bad_encoding.txt", b"\xff\xfe\xfd\xfa", "text/plain")
        response = self.client.post(
            "/api/conversions/",
            {"file": file_obj, "source_format": "txt", "target_format": "pdf"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 202)
        job = ConversionJob.objects.filter(status=JobStatus.FAILED).latest("created_at")
        if job.input_path:
            self.assertFalse(os.path.exists(job.input_path))
        if job.output_path:
            self.assertFalse(os.path.exists(job.output_path))



    def test_zero_orphaned_temporary_files(self):
        staging_dir = Path(getattr(settings, "CONVERSION_TEMP_DIR", tempfile.gettempdir()))
        initial_files = set(staging_dir.glob("velto_*")) if staging_dir.exists() else set()

        file_obj = SimpleUploadedFile("clean.txt", b"Cleanup test text", "text/plain")
        response = self.client.post(
            "/api/conversions/",
            {"file": file_obj, "source_format": "txt", "target_format": "docx"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 201)
        job_id = response.data["id"]
        job = ConversionJob.objects.get(id=job_id)

        # Input file staged for conversion must be cleaned up
        self.assertFalse(os.path.exists(job.input_path))

    def test_html_to_pdf_code_and_blockquote(self):
        html = "<blockquote><p>Quote text</p></blockquote><pre><code>x = 10</code></pre>"
        inp = self._create_temp_file(html.encode("utf-8"), ".html")
        out_pdf = self._out_path("_qb.pdf")
        HtmlToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

    def test_html_to_docx_code_and_blockquote(self):
        html = "<blockquote><p>Quote text</p></blockquote><pre><code>x = 10</code></pre>"
        inp = self._create_temp_file(html.encode("utf-8"), ".html")
        out_docx = self._out_path("_qb.docx")
        HtmlToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

    def test_markdown_tables_support(self):
        md_table = "| Col 1 | Col 2 |\n| --- | --- |\n| Val 1 | Val 2 |"
        inp = self._create_temp_file(md_table.encode("utf-8"), ".md")
        out_pdf = self._out_path("_mdtbl.pdf")
        out_docx = self._out_path("_mdtbl.docx")

        MarkdownToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

        MarkdownToDocxEngine().convert(inp, out_docx)
        validate_docx_output(out_docx)

    def test_markdown_horizontal_rule_support(self):
        md_hr = "Section 1\n\n---\n\nSection 2"
        inp = self._create_temp_file(md_hr.encode("utf-8"), ".md")
        out_pdf = self._out_path("_hr.pdf")
        MarkdownToPdfEngine().convert(inp, out_pdf)
        validate_pdf_output(out_pdf)

    def test_txt_non_utf8_rejection(self):
        inp = self._create_temp_file(b"\x80\x81\x82\x83", ".txt")
        out_pdf = self._out_path("_bad_enc.pdf")
        with self.assertRaises(ConversionError):
            TxtToPdfEngine().convert(inp, out_pdf)

    def test_html_excessive_elements_limit_rejection(self):
        from unittest.mock import patch
        html_content = "<p>Test</p>" * 20
        inp = self._create_temp_file(html_content.encode("utf-8"), ".html")
        out_pdf = self._out_path("_excessive.pdf")
        with patch("apps.conversions.engines.document_engine.MAX_HTML_ELEMENTS", 5):
            with self.assertRaises(ConversionError):
                HtmlToPdfEngine().convert(inp, out_pdf)

