"""
Phase 2 tests — PPTX → PDF conversion engine using LibreOffice.

Scenarios covered:
  1. Valid PPTX signature validation.
  2. Missing presentation.xml PPTX signature validation failure.
  3. Corrupted non-ZIP file PPTX signature validation failure.
  4. Real end-to-end PPTX → PDF conversion pipeline.
  5. Correct output filename stem preservation (e.g. presentation.pptx → presentation.pdf).
  6. Generated PDF validation with PyMuPDF (page count >= 1).
  7. Missing LibreOffice executable handling.
  8. LibreOffice conversion timeout (subprocess.TimeoutExpired).
  9. LibreOffice non-zero exit code handling.
  10. Invalid generated PDF output validation failure.
  11. API POST upload success flow.
  12. API POST upload failure flow.
  13. Download endpoint verification (HTTP 200, Content-Type: application/pdf).
  14. Anonymous session isolation (cross-session access blocked with HTTP 404).
  15. Unsupported conversion pair rejection.
"""

import io
import os
import tempfile
import unittest
from unittest.mock import patch
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from django.test import TestCase, override_settings
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APIClient

from apps.conversions.engines.base import ConversionError
from apps.conversions.engines.libreoffice import find_libreoffice_executable
from apps.conversions.engines.pptx_to_pdf import PptxToPdfEngine
from apps.conversions.engines.validators import validate_pptx_signature, validate_pdf_output
from apps.conversions.formats import is_valid_conversion
from apps.conversions.models import ConversionJob, JobStatus
from apps.conversions.services import ConversionService

# ── Isolated temp directory for PPTX tests ────────────────────────────────────
TEST_TEMP_DIR = tempfile.mkdtemp(prefix="velto_phase2_pptx_tests_")


def create_sample_pptx_bytes(slides: int = 2) -> bytes:
    """Helper to generate a valid in-memory PPTX file using python-pptx."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank layout

    for i in range(1, slides + 1):
        slide = prs.slides.add_slide(blank_layout)
        tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        tf = tx_box.text_frame
        tf.text = f"Slide {i} - VELTO Conversion PPTX to PDF Test"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@override_settings(TEMP_UPLOAD_DIR=TEST_TEMP_DIR)
class PptxToPdfEngineUnitTests(TestCase):
    """Unit tests for PPTX validation, executable discovery, and PptxToPdfEngine."""

    def setUp(self):
        self.pptx_bytes = create_sample_pptx_bytes(slides=2)

    def test_validate_pptx_signature_valid(self):
        """Valid PPTX file passes signature validation."""
        pptx_file = os.path.join(TEST_TEMP_DIR, "valid.pptx")
        with open(pptx_file, "wb") as f:
            f.write(self.pptx_bytes)
        try:
            validate_pptx_signature(pptx_file)
        finally:
            if os.path.exists(pptx_file):
                os.remove(pptx_file)

    def test_validate_pptx_signature_missing_presentation_xml(self):
        """ZIP file without ppt/presentation.xml raises ConversionError."""
        import zipfile
        bad_zip = os.path.join(TEST_TEMP_DIR, "no_ppt.pptx")
        with zipfile.ZipFile(bad_zip, "w") as zf:
            zf.writestr("word/document.xml", "<dummy/>")

        try:
            with self.assertRaises(ConversionError) as ctx:
                validate_pptx_signature(bad_zip)
            self.assertIn("valid PPTX", str(ctx.exception))
        finally:
            if os.path.exists(bad_zip):
                os.remove(bad_zip)

    def test_validate_pptx_signature_corrupted_non_zip(self):
        """Non-ZIP corrupt file raises ConversionError."""
        bad_file = os.path.join(TEST_TEMP_DIR, "corrupt.pptx")
        with open(bad_file, "wb") as f:
            f.write(b"CORRUPTED NON ZIP DATA")

        try:
            with self.assertRaises(ConversionError) as ctx:
                validate_pptx_signature(bad_file)
            self.assertIn("valid PPTX", str(ctx.exception))
        finally:
            if os.path.exists(bad_file):
                os.remove(bad_file)

    def test_missing_libreoffice_fails_cleanly(self):
        """Missing LibreOffice executable raises clear ConversionError."""
        engine = PptxToPdfEngine()
        pptx_file = os.path.join(TEST_TEMP_DIR, "test.pptx")
        with open(pptx_file, "wb") as f:
            f.write(self.pptx_bytes)

        try:
            with patch("apps.conversions.engines.libreoffice.find_libreoffice_executable", return_value=None):
                with self.assertRaises(ConversionError) as ctx:
                    engine.convert(pptx_file, os.path.join(TEST_TEMP_DIR, "out.pdf"))
                self.assertIn("LibreOffice is not installed or could not be found", str(ctx.exception))
        finally:
            if os.path.exists(pptx_file):
                os.remove(pptx_file)

    def test_libreoffice_timeout_fails_cleanly(self):
        """Subprocess timeout raises clear timeout ConversionError."""
        engine = PptxToPdfEngine()
        pptx_file = os.path.join(TEST_TEMP_DIR, "timeout.pptx")
        with open(pptx_file, "wb") as f:
            f.write(self.pptx_bytes)

        try:
            with patch("apps.conversions.engines.libreoffice.find_libreoffice_executable", return_value="soffice"), \
                 patch("subprocess.run", side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=60)):
                with self.assertRaises(ConversionError) as ctx:
                    engine.convert(pptx_file, os.path.join(TEST_TEMP_DIR, "out.pdf"))
                self.assertEqual(str(ctx.exception), "LibreOffice conversion timed out.")
        finally:
            if os.path.exists(pptx_file):
                os.remove(pptx_file)

    def test_libreoffice_nonzero_exit_code_fails_cleanly(self):
        """Subprocess non-zero exit code raises ConversionError."""
        engine = PptxToPdfEngine()
        pptx_file = os.path.join(TEST_TEMP_DIR, "nonzero.pptx")
        with open(pptx_file, "wb") as f:
            f.write(self.pptx_bytes)

        mock_res = subprocess.CompletedProcess(args=["soffice"], returncode=1, stdout="", stderr="Fatal error")

        try:
            with patch("apps.conversions.engines.libreoffice.find_libreoffice_executable", return_value="soffice"), \
                 patch("subprocess.run", return_value=mock_res):
                with self.assertRaises(ConversionError) as ctx:
                    engine.convert(pptx_file, os.path.join(TEST_TEMP_DIR, "out.pdf"))
                self.assertEqual(str(ctx.exception), "LibreOffice failed to convert the PPTX file to PDF.")
        finally:
            if os.path.exists(pptx_file):
                os.remove(pptx_file)

    def test_invalid_generated_pdf_fails_cleanly(self):
        """If LibreOffice output PDF is corrupted or invalid, validation rejects it."""
        engine = PptxToPdfEngine()
        pptx_file = os.path.join(TEST_TEMP_DIR, "badpdf.pptx")
        with open(pptx_file, "wb") as f:
            f.write(self.pptx_bytes)

        try:
            with patch("apps.conversions.engines.libreoffice.find_libreoffice_executable", return_value="soffice"), \
                 patch("apps.conversions.engines.libreoffice.validate_pdf_output", side_effect=ConversionError("The generated PDF file is invalid or unreadable.")):
                def mock_run(cmd, **kwargs):
                    out_dir = Path(cmd[cmd.index("--outdir") + 1])
                    (out_dir / "badpdf.pdf").write_bytes(b"invalid pdf bytes")
                    return subprocess.CompletedProcess(cmd, returncode=0, stdout="", stderr="")

                with patch("subprocess.run", side_effect=mock_run):
                    with self.assertRaises(ConversionError) as ctx:
                        engine.convert(pptx_file, os.path.join(TEST_TEMP_DIR, "out.pdf"))
                    self.assertEqual(str(ctx.exception), "The generated PDF file is invalid or unreadable.")
        finally:
            if os.path.exists(pptx_file):
                os.remove(pptx_file)

    def test_unsupported_conversion_pair_rejected(self):
        """Unsupported format pair (e.g. pptx -> docx) returns False."""
        self.assertFalse(is_valid_conversion("pptx", "docx"))
        self.assertTrue(is_valid_conversion("pptx", "pdf"))


@override_settings(TEMP_UPLOAD_DIR=TEST_TEMP_DIR)
class PptxToPdfApiIntegrationTests(TestCase):
    """API Integration tests for PPTX → PDF conversion flow."""

    def setUp(self):
        self.client_a = APIClient()
        self.client_b = APIClient()
        self.pptx_bytes = create_sample_pptx_bytes(slides=3)

    def test_successful_pptx_to_pdf_flow_and_download(self):
        """Full success flow: Upload PPTX → Status Completed → Download PDF."""
        exe = find_libreoffice_executable()
        if not exe:
            self.skipTest("LibreOffice not installed on machine.")

        f = io.BytesIO(self.pptx_bytes)
        f.name = "company_presentation.pptx"

        response = self.client_a.post(
            "/api/conversions/",
            {
                "file": f,
                "source_format": "pptx",
                "target_format": "pdf",
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        data = response.json()
        self.assertEqual(data["status"], "completed")
        self.assertEqual(data["source_format"], "pptx")
        self.assertEqual(data["target_format"], "pdf")
        self.assertEqual(data["output_filename"], "company_presentation.pdf")
        self.assertTrue(data["output_size_bytes"] > 0)
        self.assertIsNotNone(data["download_url"])

        job_id = data["id"]
        job = ConversionJob.objects.get(id=job_id)
        self.assertTrue(Path(job.output_path).exists())

        # Validate generated PDF page count with PyMuPDF
        import fitz
        doc = fitz.open(job.output_path)
        self.assertGreaterEqual(len(doc), 1)
        doc.close()

        # Download the PDF
        download_url = reverse("conversions:job-download", kwargs={"job_id": job_id})
        dl_response = self.client_a.get(download_url)
        self.assertEqual(dl_response.status_code, status.HTTP_200_OK)
        self.assertEqual(dl_response["Content-Type"], "application/pdf")
        content = b"".join(dl_response.streaming_content)
        self.assertTrue(content.startswith(b"%PDF"))

    def test_invalid_pptx_api_failure_flow(self):
        """Uploading corrupted PPTX file yields clean API failure response."""
        f = io.BytesIO(b"NOT A REAL PPTX")
        f.name = "corrupt.pptx"

        response = self.client_a.post(
            "/api/conversions/",
            {
                "file": f,
                "source_format": "pptx",
                "target_format": "pdf",
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
        data = response.json()
        self.assertEqual(data["status"], "failed")
        self.assertIn("valid PPTX", data["error_message"])

    def test_anonymous_session_isolation_pptx(self):
        """Job created by Session A cannot be retrieved or downloaded by Session B."""
        exe = find_libreoffice_executable()
        if not exe:
            self.skipTest("LibreOffice not installed on machine.")

        f = io.BytesIO(self.pptx_bytes)
        f.name = "confidential.pptx"

        res_a = self.client_a.post(
            "/api/conversions/",
            {
                "file": f,
                "source_format": "pptx",
                "target_format": "pdf",
            },
            format="multipart",
        )
        self.assertEqual(res_a.status_code, status.HTTP_201_CREATED)
        job_id = res_a.json()["id"]

        # Session B attempt detail view
        detail_url = reverse("conversions:job-detail", kwargs={"job_id": job_id})
        res_b_detail = self.client_b.get(detail_url)
        self.assertEqual(res_b_detail.status_code, status.HTTP_404_NOT_FOUND)

        # Session B attempt download
        dl_url = reverse("conversions:job-download", kwargs={"job_id": job_id})
        res_b_dl = self.client_b.get(dl_url)
        self.assertEqual(res_b_dl.status_code, status.HTTP_404_NOT_FOUND)
